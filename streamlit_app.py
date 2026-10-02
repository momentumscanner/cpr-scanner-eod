import streamlit as st
import pandas as pd
import io
import os
import re
from scanner import CamarillaScanner
from cpr_scanner import CPRScanner
from openpyxl.styles import Alignment, Font, PatternFill

# Page configuration
st.set_page_config(
    page_title="NSE Option Scanner (Camarilla & CPR)",
    page_icon="📈",
    layout="wide"
)

# Custom CSS
st.markdown("""
    <style>
    .main {
        padding: 2rem;
    }
    .stButton>button {
        width: 100%;
        background-color: #1f4e78;
        color: white;
        height: 3.2em;
        font-weight: bold;
        font-size: 1.1rem;
        border-radius: 6px;
    }
    </style>
""", unsafe_allow_html=True)

def generate_camarilla_excel(df, top_n=5):
    """Generates the Camarilla Excel report in memory."""
    output = io.BytesIO()
    cam_scanner = CamarillaScanner()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        cam_scanner.generate_excel_writer(df, writer, top_n=top_n)
    return output.getvalue()

def generate_cpr_excel(df, top_n=5):
    """Generates the CPR Excel report in memory."""
    output = io.BytesIO()
    cpr_scanner = CPRScanner()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        cpr_scanner.generate_excel_writer(df, writer, top_n=top_n)
    return output.getvalue()

def generate_combined_excel(df_cam, df_cpr, output_filepath_or_bytes, top_n=5):
    """Generates a combined Excel report containing all Camarilla and CPR sheets."""
    if isinstance(output_filepath_or_bytes, str):
        writer = pd.ExcelWriter(output_filepath_or_bytes, engine='openpyxl')
        should_close = True
    else:
        writer = output_filepath_or_bytes
        should_close = False

    header_font = Font(bold=True, color="FFFFFF")
    header_fill_cam = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    header_fill_cpr = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    center = Alignment(horizontal='center', vertical='center')
    cols_to_show = ['Symbol', 'Spot_Close', 'ATM_Strike']

    def add_side_by_side(df_source, filter_col, sheet_name, title_ce, title_pe, fill_style):
        df_filtered = df_source[df_source[filter_col] == True].copy()
        df_ce = df_filtered[df_filtered['Option_Type'] == 'CE'][cols_to_show].copy().reset_index(drop=True)
        df_pe = df_filtered[df_filtered['Option_Type'] == 'PE'][cols_to_show].copy().reset_index(drop=True)
        df_combined = pd.concat([df_ce, df_pe], axis=1)
        
        df_combined.to_excel(writer, sheet_name=sheet_name, index=False, startrow=1)
        worksheet = writer.sheets[sheet_name]
        
        worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=3)
        c1 = worksheet.cell(row=1, column=1, value=title_ce)
        c1.font = header_font
        c1.fill = fill_style
        c1.alignment = center
        
        worksheet.merge_cells(start_row=1, start_column=4, end_row=1, end_column=6)
        c2 = worksheet.cell(row=1, column=4, value=title_pe)
        c2.font = header_font
        c2.fill = fill_style
        c2.alignment = center

        for i, col in enumerate(df_combined.columns):
            col_idx = i + 1
            series = df_combined.iloc[:, i]
            max_len = max((series.apply(str).map(len).max() if not series.empty else 0), len(str(col))) + 2
            column_letter = worksheet.cell(row=2, column=col_idx).column_letter
            worksheet.column_dimensions[column_letter].width = max_len

    # 1. Camarilla Main Data
    df_cam_final = df_cam[['Symbol', 'Expiry', 'ATM_Strike', 'Option_Type', 'Spot_Close'] + [c for c in df_cam.columns if c not in ['Symbol', 'Expiry', 'ATM_Strike', 'Option_Type', 'Spot_Close']]]
    df_cam_final.to_excel(writer, sheet_name='Camarilla Main Data', index=False)

    # 2. CPR Main Data
    df_cpr_final = df_cpr[['Symbol', 'Expiry', 'ATM_Strike', 'Option_Type', 'Spot_Close'] + [c for c in df_cpr.columns if c not in ['Symbol', 'Expiry', 'ATM_Strike', 'Option_Type', 'Spot_Close']]]
    df_cpr_final.to_excel(writer, sheet_name='CPR Main Data', index=False)

    # Camarilla side-by-side sheets
    add_side_by_side(df_cam, 'Is_Inside_Camarilla', 'Narrow Camarilla', 'Narrow Camarilla CE', 'Narrow Camarilla PE', header_fill_cam)
    add_side_by_side(df_cam, 'Is_Inside_H4_L4', 'Inside Camarilla', 'Inside Camarilla CE', 'Inside Camarilla PE', header_fill_cam)
    add_side_by_side(df_cam, 'Is_Higher_Value', 'Higher Value Camarilla', 'Higher Value Camarilla CE', 'Higher Value Camarilla PE', header_fill_cam)
    add_side_by_side(df_cam, 'Is_Lower_Value', 'Lower Value Camarilla', 'Lower Value Camarilla CE', 'Lower Value Camarilla PE', header_fill_cam)

    # CPR side-by-side sheets
    add_side_by_side(df_cpr, 'Is_Narrow_CPR', 'Narrow CPR', 'Narrow CPR CE', 'Narrow CPR PE', header_fill_cpr)
    add_side_by_side(df_cpr, 'Is_Wide_CPR', 'Wide CPR', 'Wide CPR CE', 'Wide CPR PE', header_fill_cpr)
    add_side_by_side(df_cpr, 'Is_Inside_CPR', 'Inside CPR', 'Inside CPR CE', 'Inside CPR PE', header_fill_cpr)
    add_side_by_side(df_cpr, 'Is_Outside_CPR', 'Outside CPR', 'Outside CPR CE', 'Outside CPR PE', header_fill_cpr)
    add_side_by_side(df_cpr, 'Is_Higher_Value_CPR', 'Higher Value CPR', 'Higher Value CPR CE', 'Higher Value CPR PE', header_fill_cpr)
    add_side_by_side(df_cpr, 'Is_Lower_Value_CPR', 'Lower Value CPR', 'Lower Value CPR CE', 'Lower Value CPR PE', header_fill_cpr)

    if should_close:
        writer.close()

def render_cpr_category_tab(df, filter_col, title):
    """Renders a formatted tab for a specific CPR category."""
    df_filtered = df[df[filter_col] == True].copy()
    if df_filtered.empty:
        st.info(f"No stocks found matching **{title}** condition.")
        return
    
    unique_symbols = df_filtered['Symbol'].nunique()
    st.caption(f"Found **{unique_symbols}** symbols (**{len(df_filtered)}** total option records) for {title}.")

    cols_to_show = ['Symbol', 'Spot_Close', 'ATM_Strike', 'Today_CPR_Width', 'Today_P', 'Today_TC', 'Today_BC']
    
    col_ce, col_pe = st.columns(2)
    df_ce = df_filtered[df_filtered['Option_Type'] == 'CE'][cols_to_show].reset_index(drop=True)
    df_pe = df_filtered[df_filtered['Option_Type'] == 'PE'][cols_to_show].reset_index(drop=True)
    
    with col_ce:
        st.markdown("##### 🟢 Call Options (CE)")
        st.dataframe(df_ce, use_container_width=True)

    with col_pe:
        st.markdown("##### 🔴 Put Options (PE)")
        st.dataframe(df_pe, use_container_width=True)

def render_camarilla_category_tab(df, filter_col, title):
    """Renders a formatted tab for a specific Camarilla category."""
    df_filtered = df[df[filter_col] == True].copy()
    if df_filtered.empty:
        st.info(f"No stocks found matching **{title}** condition.")
        return
    
    unique_symbols = df_filtered['Symbol'].nunique()
    st.caption(f"Found **{unique_symbols}** symbols (**{len(df_filtered)}** total option records) for {title}.")

    cols_to_show = ['Symbol', 'Spot_Close', 'ATM_Strike']
    if 'Today_H4' in df_filtered.columns:
        cols_to_show += ['Today_H4', 'Today_H3', 'Today_L3', 'Today_L4']

    col_ce, col_pe = st.columns(2)
    df_ce = df_filtered[df_filtered['Option_Type'] == 'CE'][cols_to_show].reset_index(drop=True)
    df_pe = df_filtered[df_filtered['Option_Type'] == 'PE'][cols_to_show].reset_index(drop=True)
    
    with col_ce:
        st.markdown("##### 🟢 Call Options (CE)")
        st.dataframe(df_ce, use_container_width=True)

    with col_pe:
        st.markdown("##### 🔴 Put Options (PE)")
        st.dataframe(df_pe, use_container_width=True)

# Header
st.title("📈 NSE Option Scanner Suite")
st.markdown("Upload Today's and Yesterday's NSE FO Bhav Copy files to run Camarilla and Central Pivot Range (CPR) scans.")

# Scanner Selection
scan_mode = st.radio(
    "Select Scanner Tool:",
    options=["CPR Scanner (Central Pivot Range)", "Camarilla Scanner", "Both (Camarilla & CPR)"],
    index=0,
    horizontal=True
)

col1, col2 = st.columns(2)

with col1:
    st.subheader("Today's Data")
    today_file = st.file_uploader("Upload Today's Bhav Copy (ZIP)", type=['zip'], key='today')

with col2:
    st.subheader("Yesterday's Data")
    yest_file = st.file_uploader("Upload Yesterday's Bhav Copy (ZIP)", type=['zip'], key='yest')

# Option for Top N Results
st.markdown("### Report Settings")
top_n_choice = st.radio(
    "Select Number of Top Results to Display:",
    options=[5, 10],
    index=0,
    horizontal=True,
    help="Choose whether to see Top 5 or Top 10 results in the generated Excel report."
)

if st.button("SCAN & GENERATE REPORT"):
    if today_file is not None and yest_file is not None:
        try:
            with st.spinner('Processing files... This may take a moment.'):
                today_filename = today_file.name
                match = re.search(r"(\d{8})", today_filename)
                date_str = match.group(1) if match else "Report"

                # Define Folder Directories
                cam_dir = "CAMARILLA EOD SCANNER"
                cpr_dir = "CPR EOD SCANNER"
                combined_dir = "COMBINED EOD SCANNER"

                os.makedirs(cam_dir, exist_ok=True)
                os.makedirs(cpr_dir, exist_ok=True)
                os.makedirs(combined_dir, exist_ok=True)

                df_cpr = None
                df_cam = None

                # CPR Scanner Run
                if "CPR" in scan_mode or "Both" in scan_mode:
                    today_file.seek(0)
                    yest_file.seek(0)
                    cpr_instance = CPRScanner()
                    df_cpr = cpr_instance.process_data(today_file, yest_file)
                    
                    if df_cpr is not None and not df_cpr.empty:
                        cpr_filepath = os.path.join(cpr_dir, f"CPR Scanner {date_str}.xlsx")
                        cpr_instance.generate_excel_writer(df_cpr, cpr_filepath, top_n=top_n_choice)
                        cpr_bytes = generate_cpr_excel(df_cpr, top_n=top_n_choice)

                # Camarilla Scanner Run
                if "Camarilla" in scan_mode or "Both" in scan_mode:
                    today_file.seek(0)
                    yest_file.seek(0)
                    cam_instance = CamarillaScanner()
                    df_cam = cam_instance.process_data(today_file, yest_file)
                    
                    if df_cam is not None and not df_cam.empty:
                        cam_filepath = os.path.join(cam_dir, f"Camarilla Scanner {date_str}.xlsx")
                        cam_instance.generate_excel_writer(df_cam, cam_filepath, top_n=top_n_choice)
                        cam_bytes = generate_camarilla_excel(df_cam, top_n=top_n_choice)

                # Render CPR Results UI
                if df_cpr is not None and not df_cpr.empty:
                    st.divider()
                    st.header("🎯 Central Pivot Range (CPR) Results")
                    cpr_filepath = os.path.join(cpr_dir, f"CPR Scanner {date_str}.xlsx")
                    st.success(f"✅ CPR Scan Complete! Saved to **'{cpr_filepath}'**.")

                    st.download_button(
                        label="📥 Download CPR Excel Report",
                        data=generate_cpr_excel(df_cpr, top_n=top_n_choice),
                        file_name=f"CPR Scanner {date_str}.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        key='download_cpr'
                    )

                    tab_narrow, tab_wide, tab_inside, tab_outside, tab_higher, tab_lower, tab_main = st.tabs([
                        "🔹 Narrow CPR", "🔹 Wide CPR", "🔹 Inside CPR", "🔹 Outside CPR", "🔹 Higher Value CPR", "🔹 Lower Value CPR", "📋 Main Data"
                    ])

                    with tab_narrow: render_cpr_category_tab(df_cpr, 'Is_Narrow_CPR', 'Narrow CPR')
                    with tab_wide: render_cpr_category_tab(df_cpr, 'Is_Wide_CPR', 'Wide CPR')
                    with tab_inside: render_cpr_category_tab(df_cpr, 'Is_Inside_CPR', 'Inside CPR')
                    with tab_outside: render_cpr_category_tab(df_cpr, 'Is_Outside_CPR', 'Outside CPR')
                    with tab_higher: render_cpr_category_tab(df_cpr, 'Is_Higher_Value_CPR', 'Higher Value CPR')
                    with tab_lower: render_cpr_category_tab(df_cpr, 'Is_Lower_Value_CPR', 'Lower Value CPR')
                    with tab_main:
                        st.subheader("All Main CPR Data")
                        st.dataframe(df_cpr, use_container_width=True)

                # Render Camarilla Results UI
                if df_cam is not None and not df_cam.empty:
                    st.divider()
                    st.header("📊 Camarilla Pivot Results")
                    cam_filepath = os.path.join(cam_dir, f"Camarilla Scanner {date_str}.xlsx")
                    st.success(f"✅ Camarilla Scan Complete! Saved to **'{cam_filepath}'**.")

                    st.download_button(
                        label="📥 Download Camarilla Excel Report",
                        data=generate_camarilla_excel(df_cam, top_n=top_n_choice),
                        file_name=f"Camarilla Scanner {date_str}.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        key='download_cam'
                    )

                    tab_cam_narrow, tab_cam_inside, tab_cam_higher, tab_cam_lower, tab_cam_main = st.tabs([
                        "🔹 Narrow Camarilla", "🔹 Inside Camarilla (H4/L4)", "🔹 Higher Value Camarilla", "🔹 Lower Value Camarilla", "📋 Main Data"
                    ])

                    with tab_cam_narrow: render_camarilla_category_tab(df_cam, 'Is_Inside_Camarilla', 'Narrow Camarilla')
                    with tab_cam_inside: render_camarilla_category_tab(df_cam, 'Is_Inside_H4_L4', 'Inside Camarilla')
                    with tab_cam_higher: render_camarilla_category_tab(df_cam, 'Is_Higher_Value', 'Higher Value Camarilla')
                    with tab_cam_lower: render_camarilla_category_tab(df_cam, 'Is_Lower_Value', 'Lower Value Camarilla')
                    with tab_cam_main:
                        st.subheader("All Main Camarilla Data")
                        st.dataframe(df_cam, use_container_width=True)

                # Handle Combined File Generation if Both were run
                if df_cam is not None and df_cpr is not None:
                    st.divider()
                    st.header("⚡ Combined Report (Camarilla + CPR)")
                    combined_filepath = os.path.join(combined_dir, f"Combined Scanner {date_str}.xlsx")
                    generate_combined_excel(df_cam, df_cpr, combined_filepath, top_n=top_n_choice)

                    combined_bytes = io.BytesIO()
                    with pd.ExcelWriter(combined_bytes, engine='openpyxl') as writer:
                        generate_combined_excel(df_cam, df_cpr, writer, top_n=top_n_choice)
                    
                    st.success(f"✅ Combined Scan Complete! Saved to **'{combined_filepath}'**.")
                    st.download_button(
                        label="📥 Download Combined (Camarilla + CPR) Excel Report",
                        data=combined_bytes.getvalue(),
                        file_name=f"Combined Scanner {date_str}.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        key='download_combined'
                    )

        except Exception as e:
            st.error(f"An error occurred: {str(e)}")
            st.exception(e)
    else:
        st.warning("Please upload both ZIP files.")

# Instructions
with st.expander("How to use"):
    st.markdown("""
    1. Download the FO Bhav Copy ZIP files from NSE website for Today and Yesterday.
    2. Upload them in the respective fields above.
    3. Select CPR Scanner, Camarilla Scanner, or Both.
    4. Click 'SCAN & GENERATE REPORT'.
    5. Output files are saved in:
       - **`CAMARILLA EOD SCANNER`**
       - **`CPR EOD SCANNER`**
       - **`COMBINED EOD SCANNER`**
    """)
