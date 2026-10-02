import pandas as pd
import zipfile
import os
import datetime
from openpyxl.styles import Alignment, Font, PatternFill

class CamarillaScanner:
    def __init__(self):
        pass

    def load_bhav_copy(self, zip_path):
        """Loads the CSV from the ZIP file into a DataFrame."""
        try:
            with zipfile.ZipFile(zip_path, 'r') as z:
                # Find the first CSV file
                csv_files = [f for f in z.namelist() if f.lower().endswith('.csv')]
                if not csv_files:
                    raise ValueError(f"No CSV found in {zip_path}")
                
                # Read CSV
                df = pd.read_csv(z.open(csv_files[0]))
                
                # Standardize columns (strip whitespace)
                df.columns = df.columns.str.strip()
                
                # Strip string columns
                str_cols = ['TckrSymb', 'FinInstrmTp', 'XpryDt', 'OptnTp']
                for c in str_cols:
                    if c in df.columns:
                        df[c] = df[c].astype(str).str.strip()

                # Convert Expiry to datetime for sorting
                df['XpryDt_Date'] = pd.to_datetime(df['XpryDt'], errors='coerce')
                
                return df
        except Exception as e:
            print(f"Error loading {zip_path}: {e}")
            return None

    def calculate_camarilla(self, high, low, close):
        """Calculates Camarilla pivots."""
        r = high - low
        
        if r == 0:
            return {
                'H4': close, 'H3': close, 'H2': close, 'H1': close,
                'L1': close, 'L2': close, 'L3': close, 'L4': close
            }

        data = {}
        data['H4'] = close + (r * 1.1 / 2)
        data['H3'] = close + (r * 1.1 / 4)
        data['H2'] = close + (r * 1.1 / 6)
        data['H1'] = close + (r * 1.1 / 12)
        data['L1'] = close - (r * 1.1 / 12)
        data['L2'] = close - (r * 1.1 / 6)
        data['L3'] = close - (r * 1.1 / 4)
        data['L4'] = close - (r * 1.1 / 2)
        return data

    def get_atm_strike(self, spot_price, available_strikes):
        """Returns the strike price closest to the spot price."""
        if len(available_strikes) == 0:
            return None
        return min(available_strikes, key=lambda x: abs(x - spot_price))

    def process_data(self, today_file, yesterday_file):
        print(f"Processing Today for Camarilla: {today_file}")
        print(f"Processing Yesterday for Camarilla: {yesterday_file}")

        df_today = self.load_bhav_copy(today_file)
        df_yest = self.load_bhav_copy(yesterday_file)

        if df_today is None or df_yest is None:
            return None

        # Processing Yesterday's Data
        yest_opts = df_yest[df_yest['FinInstrmTp'].isin(['STO', 'IDO'])].copy()
        
        print("Indexing Yesterday's Camarilla data...")
        yest_lookup = {}
        for idx, row in yest_opts.iterrows():
            key = (row['TckrSymb'], float(row['StrkPric']), row['OptnTp'], row['XpryDt'])
            yest_lookup[key] = {
                'Open': row['OpnPric'],
                'High': row['HghPric'],
                'Low': row['LwPric'],
                'Close': row['ClsPric']
            }

        # Process Today's Data
        today_futs = df_today[df_today['FinInstrmTp'].isin(['STF', 'IDF'])].copy()
        today_opts = df_today[df_today['FinInstrmTp'].isin(['STO', 'IDO'])].copy()

        results = []
        symbols = today_futs['TckrSymb'].unique()
        print(f"Found {len(symbols)} underlying symbols in Futures.")

        for symbol in symbols:
            futs_sym = today_futs[today_futs['TckrSymb'] == symbol]
            if futs_sym.empty:
                continue
            
            min_expiry = futs_sym['XpryDt_Date'].min()
            nearest_fut = futs_sym[futs_sym['XpryDt_Date'] == min_expiry].iloc[0]
            
            spot_close = nearest_fut['ClsPric']
            expiry_str = nearest_fut['XpryDt']
            
            opts_sym = today_opts[
                (today_opts['TckrSymb'] == symbol) & 
                (today_opts['XpryDt'] == expiry_str)
            ]
            
            if opts_sym.empty:
                continue

            available_strikes = opts_sym['StrkPric'].astype(float).unique()
            atm_strike = self.get_atm_strike(spot_close, available_strikes)
            
            if atm_strike is None:
                continue

            for opt_type in ['CE', 'PE']:
                opt_row = opts_sym[
                    (opts_sym['StrkPric'].astype(float) == atm_strike) & 
                    (opts_sym['OptnTp'] == opt_type)
                ]
                
                if opt_row.empty:
                    continue
                
                row = opt_row.iloc[0]
                
                yest_key = (symbol, float(atm_strike), opt_type, expiry_str)
                yest_data = yest_lookup.get(yest_key)
                
                yest_levels = {}
                if yest_data:
                    yest_levels = self.calculate_camarilla(
                        yest_data['High'], yest_data['Low'], yest_data['Close']
                    )
                today_levels = self.calculate_camarilla(
                    row['HghPric'], row['LwPric'], row['ClsPric']
                )

                is_inside = False
                if yest_levels:
                    if ('H4' in today_levels and 'L4' in today_levels and 
                        'H3' in yest_levels and 'L3' in yest_levels):
                        cond1 = today_levels['H4'] < yest_levels['H3']
                        cond2 = today_levels['L4'] > yest_levels['L3']
                        if cond1 and cond2:
                            is_inside = True

                is_inside_h4_l4 = False
                if yest_levels:
                    if ('H4' in today_levels and 'L4' in today_levels and 
                        'H4' in yest_levels and 'L4' in yest_levels):
                        cond_h4 = today_levels['H4'] < yest_levels['H4']
                        cond_l4 = today_levels['L4'] > yest_levels['L4']
                        if cond_h4 and cond_l4:
                            is_inside_h4_l4 = True

                res = {
                    'Symbol': symbol,
                    'Expiry': expiry_str,
                    'Spot_Close': spot_close,
                    'ATM_Strike': atm_strike,
                    'Option_Type': opt_type,
                    'Today_Open': row['OpnPric'],
                    'Today_High': row['HghPric'],
                    'Today_Low': row['LwPric'],
                    'Today_Close': row['ClsPric'],
                    'Is_Inside_Camarilla': is_inside,
                    'Is_Inside_H4_L4': is_inside_h4_l4,
                    'Is_Higher_Value': False,
                    'Is_Lower_Value': False,
                    'OpnIntrst': row['OpnIntrst'],
                    'ChngInOpnIntrst': row['ChngInOpnIntrst'],
                    'TtlTradgVol': row['TtlTradgVol'],
                    'TtlNbOfTxsExctd': row['TtlNbOfTxsExctd']
                }

                if yest_levels:
                     if 'L4' in today_levels and 'H4' in yest_levels:
                         if today_levels['L4'] > yest_levels['H4']:
                             res['Is_Higher_Value'] = True

                if yest_levels:
                     if 'H4' in today_levels and 'L4' in yest_levels:
                         if today_levels['H4'] < yest_levels['L4']:
                             res['Is_Lower_Value'] = True
                
                for k, v in today_levels.items():
                    res[f'Today_{k}'] = round(v, 2)
                
                if yest_levels:
                    for k, v in yest_levels.items():
                        res[f'Yest_{k}'] = round(v, 2)
                
                results.append(res)

        return pd.DataFrame(results)

    def generate_excel_writer(self, df, output_writer_or_path, top_n=5):
        """Generates formatted Excel sheets for Camarilla Scanner."""
        if isinstance(output_writer_or_path, str):
            writer = pd.ExcelWriter(output_writer_or_path, engine='openpyxl')
            should_close = True
        else:
            writer = output_writer_or_path
            should_close = False

        cols = list(df.columns)
        priority = ['Symbol', 'Expiry', 'ATM_Strike', 'Option_Type', 'Spot_Close']
        final_cols = priority + [c for c in cols if c not in priority]
        df_final = df[final_cols]
        cols_to_show = ['Symbol', 'Spot_Close', 'ATM_Strike']

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
        center = Alignment(horizontal='center', vertical='center')

        def build_side_by_side_sheet(filter_col, sheet_title, ce_header_text, pe_header_text):
            df_filtered = df[df[filter_col] == True].copy()
            df_ce = df_filtered[df_filtered['Option_Type'] == 'CE'][cols_to_show].copy().reset_index(drop=True)
            df_pe = df_filtered[df_filtered['Option_Type'] == 'PE'][cols_to_show].copy().reset_index(drop=True)
            
            df_combined = pd.concat([df_ce, df_pe], axis=1)
            df_combined.to_excel(writer, sheet_name=sheet_title, index=False, startrow=1)
            
            worksheet = writer.sheets[sheet_title]
            worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=3)
            c_ce = worksheet.cell(row=1, column=1)
            c_ce.value = ce_header_text
            c_ce.font = header_font
            c_ce.fill = header_fill
            c_ce.alignment = center
            
            worksheet.merge_cells(start_row=1, start_column=4, end_row=1, end_column=6)
            c_pe = worksheet.cell(row=1, column=4)
            c_pe.value = pe_header_text
            c_pe.font = header_font
            c_pe.fill = header_fill
            c_pe.alignment = center

            for i, col in enumerate(df_combined.columns):
                col_idx = i + 1
                series = df_combined.iloc[:, i]
                max_len = max((series.apply(str).map(len).max() if not series.empty else 0), len(str(col))) + 2
                column_letter = worksheet.cell(row=2, column=col_idx).column_letter
                worksheet.column_dimensions[column_letter].width = max_len

        # Sheet 1: Main Data
        df_final.to_excel(writer, sheet_name='Main Data', index=False)

        # Sheet 2: Narrow Camarilla
        build_side_by_side_sheet('Is_Inside_Camarilla', 'Narrow Camarilla', 'Narrow Camarilla CE', 'Narrow Camarilla PE')

        # Sheet 3: Inside Camarilla
        build_side_by_side_sheet('Is_Inside_H4_L4', 'Inside Camarilla', 'Inside Camarilla CE', 'Inside Camarilla PE')

        # Sheet 4: Higher Value Camarilla
        build_side_by_side_sheet('Is_Higher_Value', 'Higher Value Camarilla', 'Higher Value Camarilla CE', 'Higher Value Camarilla PE')

        # Sheet 5: Lower Value Camarilla
        build_side_by_side_sheet('Is_Lower_Value', 'Lower Value Camarilla', 'Lower Value Camarilla CE', 'Lower Value Camarilla PE')

        # Sheet 6: Top N Output
        metrics = [
            ('OpnIntrst', f'Top {top_n} Open Interest'),
            ('ChngInOpnIntrst', f'Top {top_n} Change in OI'),
            ('TtlTradgVol', f'Top {top_n} Volume'),
            ('TtlNbOfTxsExctd', f'Top {top_n} Transactions')
        ]
        
        sheet_name_top = f'Top {top_n} Output'
        workbook = writer.book
        workbook.create_sheet(sheet_name_top)
        worksheet_top = workbook[sheet_name_top]
        
        start_row = 1
        current_col = 0
        
        for metric, title in metrics:
            df[metric] = pd.to_numeric(df[metric], errors='coerce').fillna(0)
            top_df = df.sort_values(by=metric, ascending=False).head(top_n).copy()
            cols_top = ['Symbol', 'Option_Type', 'ATM_Strike', 'Spot_Close', metric]
            top_display = top_df[cols_top].copy()

            top_display.to_excel(writer, sheet_name=sheet_name_top, index=False, startrow=start_row, startcol=current_col)
            
            op_col_start = current_col + 1
            op_col_end = current_col + len(cols_top)
            
            worksheet_top.merge_cells(start_row=1, start_column=op_col_start, end_row=1, end_column=op_col_end)
            cell_title = worksheet_top.cell(row=1, column=op_col_start)
            cell_title.value = title
            cell_title.font = header_font
            cell_title.fill = header_fill
            cell_title.alignment = center
            
            for i, col in enumerate(top_display.columns):
                col_idx = op_col_start + i
                series = top_display.iloc[:, i]
                max_len = max((series.apply(str).map(len).max() if not series.empty else 0), len(str(col))) + 2
                column_letter = worksheet_top.cell(row=2, column=col_idx).column_letter
                worksheet_top.column_dimensions[column_letter].width = max_len

            current_col += len(cols_top) + 1

        if should_close:
            writer.close()

if __name__ == "__main__":
    import glob
    files = sorted(glob.glob("BhavCopy*.zip"))
    if len(files) >= 2:
        scanner = CamarillaScanner()
        df = scanner.process_data(files[-1], files[-2])
        if df is not None:
            os.makedirs("CAMARILLA EOD SCANNER", exist_ok=True)
            output_file = os.path.join("CAMARILLA EOD SCANNER", "Camarilla_Test_Output.xlsx")
            scanner.generate_excel_writer(df, output_file)
            print(f"Saved Camarilla Test Output to {output_file}")
